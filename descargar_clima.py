import requests
import csv

from openpyxl.chart import LineChart, Reference
from openpyxl.chart.series import SeriesLabel

from datetime import datetime
from openpyxl import load_workbook

from copy import copy
from openpyxl.formula.translate import Translator

# -----------------------------------
# 1. ABRIR EXCEL Y LEER PARÁMETROS
# -----------------------------------

archivo_excel = "Sistema_Riego_Inteligente_demo.xlsx"

libro = load_workbook(archivo_excel)
hoja_parametros = libro["parametros"]


def buscar_parametro(nombre):
    for fila in hoja_parametros.iter_rows():
        for celda in fila:
            if celda.value == nombre:
                return hoja_parametros.cell(
                    row=celda.row,
                    column=celda.column + 1
                ).value

    return None


lugar = buscar_parametro("ubicacion_clima")
fecha_inicio = buscar_parametro("fecha_inicio_clima")
fecha_fin = buscar_parametro("fecha_fin_clima")


# -----------------------------------
# 2. CONVERTIR FECHAS DE EXCEL
# -----------------------------------

if isinstance(fecha_inicio, datetime):
    fecha_inicio = fecha_inicio.strftime("%Y-%m-%d")
else:
    fecha_inicio = str(fecha_inicio)

if isinstance(fecha_fin, datetime):
    fecha_fin = fecha_fin.strftime("%Y-%m-%d")
else:
    fecha_fin = str(fecha_fin)


print("Ubicación leída:", lugar)
print("Fecha inicio:", fecha_inicio)
print("Fecha fin:", fecha_fin)


# -----------------------------------
# 3. VALIDAR FECHAS
# -----------------------------------

try:
    inicio = datetime.strptime(fecha_inicio, "%Y-%m-%d")
    fin = datetime.strptime(fecha_fin, "%Y-%m-%d")

except ValueError:
    print("Error: las fechas deben tener formato AAAA-MM-DD.")
    exit()


if inicio > fin:
    print("Error: la fecha de inicio no puede ser posterior a la fecha final.")
    exit()


# -----------------------------------
# 4. BUSCAR COORDENADAS
# -----------------------------------

url_geo = "https://geocoding-api.open-meteo.com/v1/search"

parametros_geo = {
    "name": lugar,
    "count": 1,
    "language": "es",
    "countryCode": "CL"
}


try:
    respuesta_geo = requests.get(
        url_geo,
        params=parametros_geo,
        timeout=10
    )

    respuesta_geo.raise_for_status()

except requests.exceptions.RequestException:
    print("Error: no se pudo conectar con el servicio de ubicación.")
    exit()


datos_geo = respuesta_geo.json()


if "results" not in datos_geo:
    print("No se encontró la ubicación.")
    exit()


resultado = datos_geo["results"][0]

latitud = resultado["latitude"]
longitud = resultado["longitude"]


print()
print("Lugar encontrado:", resultado["name"])
print("Región:", resultado.get("admin1", "No disponible"))
print("País:", resultado.get("country", "No disponible"))
print("Latitud:", latitud)
print("Longitud:", longitud)


# -----------------------------------
# 5. CONSULTAR DATOS CLIMÁTICOS
# -----------------------------------

url = (
    f"https://archive-api.open-meteo.com/v1/archive"
    f"?latitude={latitud}"
    f"&longitude={longitud}"
    f"&start_date={fecha_inicio}"
    f"&end_date={fecha_fin}"
    f"&daily=et0_fao_evapotranspiration,precipitation_sum"
    f"&timezone=auto"
)


try:
    respuesta = requests.get(url, timeout=10)
    respuesta.raise_for_status()

except requests.exceptions.RequestException:
    print("Error: no se pudo conectar con el servicio meteorológico.")
    exit()


datos = respuesta.json()
diario = datos["daily"]

fechas = diario["time"]
eto = diario["et0_fao_evapotranspiration"]
precipitacion = diario["precipitation_sum"]


# -----------------------------------
# 6. CREAR CSV
# -----------------------------------

nombre_lugar = resultado["name"].replace(" ", "_")

nombre_archivo = (
    f"clima_{nombre_lugar}_{fecha_inicio}_{fecha_fin}.csv"
)


with open(
    nombre_archivo,
    "w",
    newline="",
    encoding="utf-8"
) as archivo:

    escritor = csv.writer(archivo, delimiter=";")

    escritor.writerow([
        "fecha",
        "eto_mm",
        "precipitacion_mm",
        "fuente"
    ])

    for fecha, eto_dia, precipitacion_dia in zip(
        fechas,
        eto,
        precipitacion
    ):

        eto_csv = str(eto_dia).replace(".", ",")
        precipitacion_csv = str(precipitacion_dia).replace(".", ",")

        escritor.writerow([
            fecha,
            eto_csv,
            precipitacion_csv,
            "Open-Meteo"
        ])


print()
print("Archivo creado:", nombre_archivo)


# -----------------------------------
# 7. ACTUALIZAR HOJA clima_real
# -----------------------------------

hoja_clima = libro["clima_real"]


hoja_clima["A1"] = "fecha"
hoja_clima["B1"] = "eto_mm"
hoja_clima["C1"] = "precipitacion_mm"
hoja_clima["D1"] = "fuente"
hoja_clima["E1"] = "precipitacion_efectiva_mm"


# Borrar datos anteriores
for fila in range(2, hoja_clima.max_row + 1):
    for columna in range(1, 6):
        hoja_clima.cell(
            row=fila,
            column=columna
        ).value = None


# Escribir datos nuevos
for fila, (fecha, eto_dia, precipitacion_dia) in enumerate(
    zip(fechas, eto, precipitacion),
    start=2
):

    fecha_excel = datetime.strptime(
        fecha,
        "%Y-%m-%d"
    )

    hoja_clima.cell(
        row=fila,
        column=1
    ).value = fecha_excel

    hoja_clima.cell(
        row=fila,
        column=1
    ).number_format = "dd/mm/yyyy"

    hoja_clima.cell(
        row=fila,
        column=2
    ).value = eto_dia

    hoja_clima.cell(
        row=fila,
        column=3
    ).value = precipitacion_dia

    hoja_clima.cell(
        row=fila,
        column=4
    ).value = "Open-Meteo"

    # Simplificación V2
    hoja_clima.cell(
        row=fila,
        column=5
    ).value = precipitacion_dia


# -----------------------------------
# 8. GUARDAR EXCEL
# -----------------------------------

# -----------------------------------
# 8. AJUSTAR riego_diario AL RANGO CLIMÁTICO
# -----------------------------------

hoja_riego = libro["riego_diario"]

# La fila 3 sirve como plantilla para todos los días
# posteriores al primer día.
plantilla_fila = 3

# Guardamos la plantilla antes de modificar las filas
plantilla = []

for columna in range(2, 32):  # B hasta AE
    celda = hoja_riego.cell(row=plantilla_fila, column=columna)

    plantilla.append({
        "columna": columna,
        "valor": celda.value,
        "estilo": copy(celda._style)
    })


# Número de filas necesarias según los datos climáticos
cantidad_dias = len(fechas)
ultima_fila_nueva = cantidad_dias + 1


# Primer día: solo actualizamos la fecha.
# La fila 2 conserva sus fórmulas especiales.
fecha_primera = datetime.strptime(fechas[0], "%Y-%m-%d")

hoja_riego["A2"] = fecha_primera
hoja_riego["A2"].number_format = "dd/mm/yyyy"


# Días siguientes
for indice, fecha in enumerate(fechas[1:], start=3):

    fecha_excel = datetime.strptime(fecha, "%Y-%m-%d")

    hoja_riego.cell(row=indice, column=1).value = fecha_excel
    hoja_riego.cell(row=indice, column=1).number_format = "dd/mm/yyyy"

    for dato in plantilla:

        columna = dato["columna"]
        valor = dato["valor"]

        celda_destino = hoja_riego.cell(
            row=indice,
            column=columna
        )

        celda_destino._style = copy(dato["estilo"])

        # Si es fórmula, adaptar referencias a la nueva fila
        if isinstance(valor, str) and valor.startswith("="):

            origen = hoja_riego.cell(
                row=plantilla_fila,
                column=columna
            ).coordinate

            destino = celda_destino.coordinate

            celda_destino.value = Translator(
                valor,
                origin=origen
            ).translate_formula(destino)

        else:
            # Copiar datos constantes:
            # cultivo, etapa, superficie, eficiencia, caudal...
            celda_destino.value = valor


# Borrar filas antiguas que hayan quedado fuera del nuevo rango
for fila in range(
    ultima_fila_nueva + 1,
    hoja_riego.max_row + 1
):
    for columna in range(1, 32):
        hoja_riego.cell(
            row=fila,
            column=columna
        ).value = None

# =========================================================
# CREAR GRAFICO DEL PANEL
# =========================================================

hoja_panel = libro["panel"]

# Eliminar gráficos anteriores para evitar duplicados
hoja_panel._charts = []

grafico = LineChart()

grafico.title = "Déficit hídrico vs. umbral RAW"
grafico.y_axis.title = "Agua (mm)"
grafico.x_axis.title = "Fecha"

grafico.height = 8
grafico.width = 18

# Déficit acumulado: columna P
datos_deficit = Reference(
    hoja_riego,
    min_col=16,
    min_row=1,
    max_row=ultima_fila_nueva
)

# RAW: columna R
datos_raw = Reference(
    hoja_riego,
    min_col=18,
    min_row=1,
    max_row=ultima_fila_nueva
)

# Fechas: columna A
fechas_grafico = Reference(
    hoja_riego,
    min_col=1,
    min_row=2,
    max_row=ultima_fila_nueva
)

grafico.add_data(datos_deficit, titles_from_data=True)
grafico.add_data(datos_raw, titles_from_data=True)
grafico.set_categories(fechas_grafico)

# Nombres limpios en la leyenda
grafico.series[0].tx = SeriesLabel(v="Déficit acumulado")
grafico.series[1].tx = SeriesLabel(v="Umbral RAW")

grafico.legend.position = "b"

# Posición dentro del panel
hoja_panel.add_chart(grafico, "A13")

libro.save(archivo_excel)

print("Excel actualizado:", archivo_excel)